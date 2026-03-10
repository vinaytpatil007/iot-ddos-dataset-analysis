from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE_DIR = Path(r"D:\ddos dataset")
ANALYSIS_DIR = BASE_DIR / "analysis_results"
ADVANCED_DIR = ANALYSIS_DIR / "advanced_analysis"
OUTPUT_PATH = BASE_DIR / "iot_ddos_dataset_analysis_paper_v5_final.docx"

TITLE = "Comparative Analysis of Modern IoT DDoS Datasets: CICIoT2023, CICDDoS2019, and TON-IoT"
AUTHOR_NAME = "Vinay T. Patil"
AFFILIATION_1 = "Department of Computer Engineering"
AFFILIATION_2 = "Kavayitri Bahinabai Chaudhari North Maharashtra University"
AUTHOR_EMAIL = "Email: vinay@example.com"


def load_tables() -> dict:
    dataset_table = pd.read_csv(ANALYSIS_DIR / "dataset_comparison_table.csv")
    dataset_table["Year"] = dataset_table["Dataset"].map(
        {
            "CICIoT2023": 2023,
            "CICDDoS2019": 2019,
            "TONIoT": 2020,
            "BoTIoT_optional": 2018,
        }
    )
    dataset_table = dataset_table[
        ["Dataset", "Year", "Total Samples", "DDoS Samples", "Benign Samples", "Features", "Class Imbalance Ratio"]
    ]
    return {
        "dataset_comparison": dataset_table,
        "dataset_difficulty": pd.read_csv(ANALYSIS_DIR / "dataset_difficulty_table.csv"),
        "feature_redundancy": pd.read_csv(ANALYSIS_DIR / "feature_redundancy_summary.csv"),
    }


def set_default_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(11)


def set_margins(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_title_block(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in [AUTHOR_NAME, AFFILIATION_1, AFFILIATION_2, AUTHOR_EMAIL]:
        block = document.add_paragraph()
        block.alignment = WD_ALIGN_PARAGRAPH.CENTER
        block.add_run(line)

    document.add_paragraph("")


def add_heading(document: Document, text: str, level: int = 1) -> None:
    p = document.add_paragraph()
    p.style = document.styles[f"Heading {level}"]
    run = p.add_run(text)
    run.bold = True


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.15
        p.add_run(item)


def add_keywords(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    lead = p.add_run("Keywords: ")
    lead.bold = True
    p.add_run(text)


def add_table(document: Document, dataframe: pd.DataFrame, caption: str) -> None:
    cp = document.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = cp.add_run(caption)
    cap.bold = True

    table = document.add_table(rows=1, cols=len(dataframe.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = table.rows[0].cells
    for i, col in enumerate(dataframe.columns):
        headers[i].text = str(col)
        shade_cell(headers[i], "D9EAF7")
        for p in headers[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True

    for _, row in dataframe.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if isinstance(value, float):
                text = "inf" if value == float("inf") else f"{value:.6f}" if abs(value) < 1000 else f"{value:.2f}"
            else:
                text = str(value)
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph("")


def add_figure(document: Document, image_path: Path, caption: str, width: float = 5.8) -> None:
    if not image_path.exists():
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))

    cp = document.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = cp.add_run(caption)
    cap.italic = True


def references() -> list[str]:
    return [
        "[1] E. C. P. Neto, S. Dadkhah, R. Ferreira, A. Zohourian, R. Lu, and A. A. Ghorbani, \"CICIoT2023: A real-time dataset and benchmark for large-scale attacks in IoT environment,\" Sensors, vol. 23, no. 13, Art. no. 5941, 2023.",
        "[2] I. Sharafaldin, A. H. Lashkari, S. Hakak, and A. A. Ghorbani, \"Developing realistic distributed denial of service (DDoS) attack dataset and taxonomy,\" in Proc. IEEE 53rd Int. Carnahan Conf. Security Technology (ICCST), 2019, pp. 1-8.",
        "[3] A. Alsaedi, N. Moustafa, Z. Tari, A. Mahmood, and A. Anwar, \"TON_IoT telemetry dataset: A new generation dataset of IoT and IIoT for data-driven intrusion detection systems,\" IEEE Access, vol. 8, pp. 165130-165150, 2020.",
        "[4] N. Moustafa, \"A new distributed architecture for evaluating AI-based security systems at the edge: Network TON_IoT datasets,\" Sustainable Cities and Society, vol. 72, Art. no. 102994, 2021.",
        "[5] T. M. Booij, I. Chiscop, E. Meeuwissen, N. Moustafa, and F. T. H. den Hartog, \"ToN_IoT: The role of heterogeneity and the need for standardization of features and attack types in IoT intrusion datasets,\" IEEE Internet of Things Journal, vol. 9, no. 1, pp. 485-496, 2022.",
        "[6] I. Sharafaldin, A. H. Lashkari, and A. A. Ghorbani, \"Toward generating a new intrusion detection dataset and intrusion traffic characterization,\" in Proc. 4th Int. Conf. Information Systems Security and Privacy (ICISSP), 2018, pp. 108-116.",
        "[7] N. Moustafa and J. Slay, \"UNSW-NB15: A comprehensive data set for network intrusion detection systems,\" in Proc. Military Communications and Information Systems Conf. (MilCIS), 2015, pp. 1-6.",
        "[8] N. Moustafa and J. Slay, \"The evaluation of network anomaly detection systems: Statistical analysis of the UNSW-NB15 data set and the comparison with the KDD99 data set,\" Information Security Journal: A Global Perspective, vol. 25, no. 1-3, pp. 18-31, 2016.",
        "[9] A. Khraisat, I. Gondal, P. Vamplew, and J. Kamruzzaman, \"Survey of intrusion detection systems: Techniques, datasets and challenges,\" Cybersecurity, vol. 2, Art. no. 20, 2019.",
        "[10] A. Adnan, A. Muhammed, A. A. A. Ghani, A. Abdullah, and F. Hakim, \"An intrusion detection system for the Internet of Things based on machine learning: Review and challenges,\" Symmetry, vol. 13, no. 6, Art. no. 1011, 2021.",
        "[11] C. Kolias, G. Kambourakis, A. Stavrou, and J. Voas, \"DDoS in the IoT: Mirai and other botnets,\" Computer, vol. 50, no. 7, pp. 80-84, 2017.",
        "[12] S. Garcia, A. Parmisano, and M. J. Erquiaga, \"IoT-23: A labeled dataset with malicious and benign IoT network traffic,\" Zenodo, 2020.",
        "[13] KDD Cup 1999 Data, UCI KDD Archive, Univ. California, Irvine, 1999.",
        "[14] M. Tavallaee, E. Bagheri, W. Lu, and A. A. Ghorbani, \"A detailed analysis of the KDD CUP 99 data set,\" in Proc. IEEE Symp. Computational Intelligence for Security and Defense Applications, 2009, pp. 1-6.",
        "[15] H. Hotelling, \"Analysis of a complex of statistical variables into principal components,\" Journal of Educational Psychology, vol. 24, no. 6, pp. 417-441, 1933.",
        "[16] L. van der Maaten and G. Hinton, \"Visualizing data using t-SNE,\" Journal of Machine Learning Research, vol. 9, pp. 2579-2605, 2008.",
        "[17] L. Breiman, \"Random forests,\" Machine Learning, vol. 45, no. 1, pp. 5-32, 2001.",
        "[18] D. R. Cox, \"The regression analysis of binary sequences,\" Journal of the Royal Statistical Society: Series B, vol. 20, no. 2, pp. 215-242, 1958.",
        "[19] B. R. Kikissagbe and M. Adda, \"Machine learning-based intrusion detection methods in IoT systems: A comprehensive review,\" Electronics, vol. 13, no. 18, Art. no. 3601, 2024.",
        "[20] N. Nisha, N. S. Gill, and P. Gulia, \"A review on machine learning based intrusion detection system for internet of things enabled environment,\" International Journal of Electrical and Computer Engineering, vol. 14, no. 2, pp. 1890-1898, 2024.",
        "[21] A. Pinto, L.-C. Herrera, Y. Donoso, and J. A. Gutierrez, \"Survey on intrusion detection systems based on machine learning techniques for the protection of critical infrastructure,\" Sensors, vol. 23, no. 5, Art. no. 2415, 2023.",
        "[22] A. Awajan, \"A novel deep learning-based intrusion detection system for IoT networks,\" Computers, vol. 12, no. 2, Art. no. 34, 2023.",
        "[23] Z. Wang, H. Chen, S. Yang, X. Luo, D. Li, and J. Wang, \"A lightweight intrusion detection method for IoT based on deep learning and dynamic quantization,\" PeerJ Computer Science, vol. 9, Art. no. e1569, 2023.",
        "[24] A. U. R. Khan, \"A survey of deep learning technologies for intrusion detection in Internet of Things,\" IEEE Access, vol. 12, pp. 4745-4761, 2024.",
        "[25] M. S. Mohammed and H. A. Talib, \"Using machine learning algorithms in intrusion detection systems: A review,\" Tikrit Journal of Pure Science, vol. 29, no. 3, pp. 63-74, 2024.",
    ]


def abstract_text() -> str:
    return (
        "The rapid growth of Internet of Things (IoT) deployments has intensified the need for reliable distributed denial-of-service (DDoS) detection benchmarks and rigorous dataset-centric evaluation. This study comparatively analyzes CICIoT2023, CICDDoS2019, and TON-IoT using statistical profiling, feature redundancy analysis, dimensionality-reduction visualization, and baseline machine-learning experiments. The results show substantial differences in class balance, feature structure, and effective learning difficulty. Logistic Regression achieved ROC-AUC values of 0.999994 on CICIoT2023, 0.999212 on CICDDoS2019, and 0.932677 on TON-IoT, while Random Forest achieved 1.000000, 0.999995, and 0.999986, respectively. These findings indicate that CICIoT2023 and CICDDoS2019 are highly separable under baseline conditions, whereas TON-IoT provides a comparatively more challenging linear classification setting despite strong nonlinear separability. The study also identifies severe class imbalance in CICDDoS2019 and strong redundancy among multiple flow-oriented features. Overall, the paper demonstrates that benchmark properties materially influence reported model performance and should therefore be analyzed explicitly when selecting datasets for IoT DDoS detection research."
    )


def intro_paragraphs() -> list[str]:
    return [
        "The expansion of IoT technologies has transformed contemporary digital ecosystems by embedding sensing, communication, and control functions into homes, hospitals, factories, transportation systems, and urban infrastructure [9], [10], [19]-[25]. These deployments increasingly operate as cyber-physical systems in which digital communications and physical processes are tightly coupled, making service disruption a security issue with direct operational consequences.",
        "IoT-enabled services support smart cities, connected healthcare, industrial monitoring, environmental sensing, logistics automation, and edge-assisted decision-making. However, the same characteristics that enable wide adoption, including low-cost hardware, limited native defenses, persistent connectivity, and device heterogeneity, also expose IoT environments to large-scale exploitation and coordinated attack campaigns [10], [11], [19], [20].",
        "DDoS attacks remain one of the most disruptive threats in this space because compromised IoT devices can be rapidly transformed into botnets that exhaust bandwidth, overload services, and degrade availability at scale. The Mirai botnet demonstrated the practical consequences of this phenomenon and established IoT botnet-driven disruption as a major cybersecurity concern [11].",
        "Traditional IDS approaches are often insufficient in IoT environments because signature-based detection struggles against evolving attack variants, while rigid rule-based systems scale poorly across heterogeneous device populations and traffic profiles [9], [10], [21], [24]. Machine-learning-based IDS techniques are therefore widely studied as a means of modeling complex network behavior and adapting to broader attack patterns [19]-[25].",
        "Benchmark datasets play a decisive role in this process because they define the attack classes, class distributions, feature spaces, and traffic realism from which empirical conclusions are drawn [1]-[8], [12]-[14]. A detector evaluated on an easier or heavily imbalanced benchmark may appear highly effective without necessarily generalizing to more heterogeneous environments. This paper therefore adopts a dataset-centric perspective and comparatively analyzes CICIoT2023, CICDDoS2019, and TON-IoT in order to support more rigorous benchmark selection and interpretation."
    ]


def contributions() -> list[str]:
    return [
        "A comparative statistical analysis of three modern IoT DDoS datasets, namely CICIoT2023, CICDDoS2019, and TON-IoT.",
        "A unified characterization of dataset scale, feature count, class imbalance, and benchmark difficulty using the same analytical pipeline.",
        "A feature-centric assessment combining redundancy analysis, entropy-oriented interpretation, and distribution-aware inspection of cleaned feature spaces.",
        "A visualization-based analysis using PCA and t-SNE to examine benign-versus-DDoS separability across the evaluated datasets.",
        "A baseline performance comparison using Logistic Regression and Random Forest to estimate practical dataset difficulty under linear and nonlinear decision boundaries.",
        "A journal-oriented synthesis of dataset strengths, weaknesses, and limitations relevant to future IoT DDoS detection research."
    ]


def related_work_paragraphs() -> list[str]:
    return [
        "Intrusion detection dataset development has evolved from early enterprise-oriented corpora to more realistic and domain-specific benchmarks. Historical resources such as KDD Cup 1999 played an influential role in intrusion detection research but were later criticized for outdated attack distributions, duplicated records, and limited realism [13], [14].",
        "Modern non-IoT datasets such as UNSW-NB15 and CICIDS2017 improved this situation by incorporating more realistic traffic generation, more recent attacks, and richer flow-based feature extraction [6]-[8]. These datasets remain useful for general intrusion detection research, but they do not fully capture the heterogeneity and cyber-physical constraints of IoT ecosystems.",
        "IoT-focused benchmarks subsequently emerged to address that gap. TON-IoT introduced multi-source telemetry, operating-system traces, and network traffic from IoT and IIoT settings [3]-[5], while IoT-23 contributed labeled benign and malicious IoT traffic with realistic scenarios [12]. CICIoT2023 more recently provided a large-scale IoT benchmark explicitly designed around extensive attack coverage and real-device collection [1].",
        "In parallel, the methodological literature has increasingly emphasized machine learning and deep learning for IoT intrusion detection, including linear models, tree ensembles, deep neural networks, and lightweight architectures tailored to resource-constrained deployments [9], [10], [19]-[25]. Survey work has also highlighted persistent issues involving benchmark heterogeneity, feature standardization, and uneven evaluation practices [5], [9], [19], [21], [24].",
        "Despite these advances, comparatively fewer studies focus on the datasets themselves as research objects. This remains a critical omission because differences in dataset balance, redundancy, feature richness, and traffic realism can materially influence reported model performance. The present study addresses that gap through a structured comparative analysis of three contemporary IoT DDoS datasets."
    ]


def methodology_paragraphs() -> list[str]:
    return [
        "The methodological workflow began with the acquisition of processed benchmark subsets derived from CICIoT2023, CICDDoS2019, and TON-IoT [1]-[5]. Each subset had already been filtered to retain benign and DDoS traffic relevant to a binary detection setting. This enabled consistent comparative analysis while preserving the original benchmark identities.",
        "A data preprocessing pipeline was then applied to clean the feature spaces. Non-numeric identifiers such as IP addresses and timestamps were removed, labels were standardized into benign and DDoS classes, and remaining variables were converted into numeric form where appropriate. Because the three datasets differ substantially in schema, preprocessing prioritized clean per-dataset analysis rather than imposing an artificial universal schema.",
        "Feature extraction in this study therefore refers to the cleaned numeric attributes retained after preprocessing. Statistical analysis was used to quantify class imbalance, feature counts, skewness, and other descriptive properties. Redundancy detection relied on correlation analysis among cleaned features, while entropy analysis was used to assess feature variability and informational richness.",
        "For visualization, principal component analysis (PCA) and t-distributed stochastic neighbor embedding (t-SNE) were applied to sampled and standardized feature matrices [15], [16]. PCA provided a linear view of dominant variance structure, while t-SNE provided a nonlinear view of local clustering behavior. Extremely large datasets were sampled prior to visualization in order to keep the analysis computationally tractable.",
        "Baseline machine-learning experiments were performed using Logistic Regression and Random Forest classifiers to estimate benchmark difficulty under linear and nonlinear decision boundaries [17], [18]. Stratified splitting preserved class proportions, and performance was summarized using ROC-AUC along with standard classification metrics. These experiments were intended as dataset diagnostics rather than as optimized state-of-the-art detectors."
    ]


def dataset_description_intro() -> str:
    return (
        "The three datasets analyzed in this study reflect distinct benchmark design goals. CICIoT2023 emphasizes large-scale IoT attack realism, CICDDoS2019 focuses on realistic DDoS taxonomy and flow-based evaluation, and TON-IoT emphasizes heterogeneity across telemetry, operating-system, and network sources [1]-[5]. Table 1 summarizes the comparative statistics used throughout this study."
    )


def ciciot_paragraphs() -> list[str]:
    return [
        "CICIoT2023 was introduced as a large-scale benchmark for IoT attacks collected in an extensive topology composed of 105 real IoT devices [1]. The documentation emphasizes that the dataset was designed to address the scarcity of realistic large-device IoT attack environments in prior literature and to support the development of operational IoT security analytics.",
        "The benchmark includes 33 attacks grouped into seven high-level classes: DDoS, DoS, Recon, Web-based, Brute Force, Spoofing, and Mirai [1]. In the DDoS-focused subset used in this work, the dataset includes multiple attack families such as SYN flood, TCP flood, UDP flood, ICMP flood, fragmentation-based attacks, RST/FIN flood, PSH-ACK flood, SlowLoris, and HTTP flood. This breadth makes the dataset especially useful for studying variation within the DDoS class.",
        "The local documentation also describes the CIC IoT laboratory as a dedicated infrastructure for realistic IoT security data collection. Raw packet captures, extracted CSV features, notebooks, and supplementary wrangling materials are included in the full distribution. In the cleaned subset analyzed here, the feature space contains 39 numeric variables derived from flow-oriented network behavior."
    ]


def cicddos_paragraphs() -> list[str]:
    return [
        "CICDDoS2019 was developed to provide both a realistic DDoS benchmark and a taxonomy-oriented view of contemporary DDoS attacks [2]. The documentation explicitly distinguishes reflection-based and exploitation-based attacks, thereby grounding the dataset in a structured family-level interpretation of DDoS behavior.",
        "The benchmark contains benign traffic and modern attacks such as DNS, LDAP, MSSQL, NetBIOS, SNMP, SSDP, UDP, UDP-Lag, SYN, NTP, TFTP, PortMap, and WebDDoS [2]. Benign traffic was generated using a user-behavior profiling approach intended to approximate realistic background activity across common protocols, thereby improving the realism of the benchmark relative to older datasets.",
        "Feature extraction was performed using CICFlowMeter-V3, with the documentation reporting more than 80 traffic features in the resulting CSV files [2]. This is consistent with the 82 cleaned numeric features retained in the present study. The resulting benchmark remains one of the richest flow-based DDoS datasets among those considered here."
    ]


def toniot_paragraphs() -> list[str]:
    return [
        "TON-IoT was designed as a new generation IoT and IIoT benchmark for AI-based cybersecurity applications and was collected in a realistic distributed testbed deployed at the UNSW Canberra Cyber Range and IoT Labs [3]-[5]. The documentation emphasizes the integration of IoT, cloud, and edge/fog layers and the use of multiple virtual machines and hosts across Windows, Linux, and Kali environments.",
        "The benchmark is heterogeneous by design. It includes telemetry from IoT and IIoT sensors, network traffic, Linux traces, Windows traces, processed feature sets, train-test subsets, and security event ground truth [3]-[5]. Attack scenarios include DoS, DDoS, ransomware, and related malicious activity against IoT gateways, web applications, MQTT-related services, and operating-system resources.",
        "This heterogeneity differentiates TON-IoT from purely flow-oriented datasets. Although the cleaned binary subset used in this study contains 19 numeric features, the broader benchmark is much richer in modality and context. This characteristic is one reason TON-IoT is analytically important: it better reflects the complexity of heterogeneous IoT and IIoT environments than simpler single-source datasets [3]-[5]."
    ]


def dataset_analysis_paragraphs() -> list[str]:
    return [
        "The comparative statistics reveal pronounced differences in benchmark composition. As illustrated in Figure 1, CICIoT2023 contributes the largest processed sample size, whereas CICDDoS2019 provides the richest cleaned numeric feature space. TON-IoT is smaller but comparatively less skewed, which makes it useful for analyzing whether strong detector performance persists under less extreme class conditions.",
        "Figure 2 highlights class imbalance across the evaluated datasets. CICDDoS2019 is the most imbalanced benchmark, CICIoT2023 is also strongly attack-dominated, and TON-IoT presents the most moderate DDoS-to-benign ratio among the three main datasets. These structural differences should be considered whenever cross-paper comparisons are made, because class imbalance materially affects the interpretation of benchmark results [9], [10], [21]."
    ]


def feature_analysis_paragraphs() -> list[str]:
    return [
        "Feature-level analysis indicates that several variables are highly skewed and multiple feature pairs are strongly redundant. This pattern is expected in engineered flow-based datasets where related packet counts, lengths, and protocol-flag measurements often encode overlapping information.",
        "Table 2 summarizes representative highly redundant feature pairs. In practical terms, such redundancy implies that reported model performance may depend on repeated signals rather than on truly diverse informational content. Consequently, feature selection and ablation analyses are well justified for future IoT DDoS experiments [8]-[10], [19], [24]."
    ]


def results_paragraphs() -> list[str]:
    return [
        "The visualization results support the statistical findings. As shown in Figures 3, 4, and 5, PCA reveals substantial benign-versus-DDoS separation for the three datasets, with the CIC benchmarks exhibiting particularly strong structure. The corresponding t-SNE projections in Figures 6, 7, and 8 reinforce this observation while also showing that TON-IoT exhibits comparatively more complex local clustering behavior.",
        "Table 3 and Figure 9 summarize the baseline difficulty comparison. Logistic Regression achieved ROC-AUC values of 0.999994 on CICIoT2023, 0.999212 on CICDDoS2019, and 0.932677 on TON-IoT, whereas Random Forest achieved 1.000000, 0.999995, and 0.999986, respectively. These results indicate that CICIoT2023 and CICDDoS2019 are highly separable under baseline conditions, while TON-IoT provides the most challenging linear classification setting in this study."
    ]


def discussion_paragraphs() -> list[str]:
    return [
        "The results indicate that benchmark difficulty is not uniform across modern IoT DDoS datasets. Near-perfect baseline performance on CICIoT2023 and CICDDoS2019 should not be interpreted as evidence that a detector will generalize equally well to more heterogeneous or less linearly separable conditions. This is one reason why TON-IoT is analytically valuable despite its smaller cleaned feature space.",
        "At the same time, the strong performance of nonlinear models across all three datasets indicates that modern engineered features remain highly informative for benchmark-driven DDoS classification. The key implication is not that these datasets are unusable, but rather that published model claims should be contextualized by dataset structure, class balance, and redundancy rather than judged from accuracy metrics alone."
    ]


def limitations_paragraphs() -> list[str]:
    return [
        "A first limitation concerns class imbalance. CICDDoS2019 is heavily dominated by DDoS traffic in the cleaned subset, which can make very strong scores easier to achieve and reduce confidence that the benign class is adequately represented. Although CICIoT2023 is less extreme, it also remains attack-heavy relative to TON-IoT.",
        "A second limitation concerns the high separability observed in the CIC benchmarks. Extremely strong ROC-AUC values are useful for benchmarking but may indicate that the evaluated subsets are easier than real operational conditions. In other words, benchmark difficulty in these datasets may underestimate the complexity of real-world IoT attack detection.",
        "A third limitation concerns redundancy among flow-oriented features. Strong correlations among multiple engineered variables indicate that some benchmarks may offer fewer truly independent signals than their raw feature counts suggest. This can distort both feature-importance interpretation and the apparent diversity of the learning problem.",
        "A fourth limitation involves real-world traffic diversity. Even carefully designed benchmarks remain abstractions of operational environments and may not capture the full range of benign variability found in heterogeneous IoT ecosystems, especially across smart city, industrial, and healthcare deployments.",
        "Finally, restricted generalization across heterogeneous IoT environments remains an open concern. Differences in device populations, telemetry modalities, network protocols, and attack implementations can create domain shift between datasets. For that reason, strong within-dataset performance should not be assumed to translate directly into cross-dataset robustness."
    ]


def conclusion_paragraphs() -> list[str]:
    return [
        "This paper presented a polished comparative analysis of CICIoT2023, CICDDoS2019, and TON-IoT for IoT DDoS detection research. By combining dataset documentation, descriptive statistics, feature analysis, visualization, and baseline modeling, the study provided a dataset-centric view of benchmark behavior rather than a purely model-centric comparison.",
        "The results confirm that the three benchmarks differ substantially in sample size, class balance, feature structure, redundancy, and practical learning difficulty. These findings reinforce the importance of transparent benchmark selection and support more rigorous interpretation of machine-learning-based IoT DDoS detection results."
    ]


def future_work_paragraphs() -> list[str]:
    return [
        "Future research should evaluate cross-dataset generalization, domain adaptation, and transfer learning so that the robustness of detectors can be assessed beyond single-benchmark settings. This is especially relevant in IoT security, where traffic heterogeneity can substantially affect model transferability.",
        "Additional directions include deep learning, sequence-aware architectures, federated evaluation pipelines, and the construction of new real-world IoT datasets with richer benign diversity and contemporary attack behavior. Such steps would help bridge the gap between high-performing benchmark experiments and operational deployment requirements."
    ]


def main() -> None:
    data = load_tables()
    dataset_table = data["dataset_comparison"]
    difficulty_table = data["dataset_difficulty"]
    redundancy_table = data["feature_redundancy"].head(12)

    document = Document()
    set_default_font(document)
    set_margins(document)
    add_page_number(document.sections[0])

    add_title_block(document)

    add_heading(document, "Abstract")
    add_body(document, abstract_text())

    add_heading(document, "Keywords")
    add_keywords(document, "IoT security, DDoS detection, dataset analysis, intrusion detection systems, machine learning, network security, cybersecurity datasets, IoT cybersecurity")

    add_heading(document, "Introduction")
    for paragraph in intro_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Research Contributions", level=2)
    add_bullets(document, contributions())

    add_heading(document, "Related Work")
    for paragraph in related_work_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Methodology")
    for paragraph in methodology_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Dataset Description")
    add_body(document, dataset_description_intro())
    add_heading(document, "4.1 CICIoT2023 Dataset", level=2)
    for paragraph in ciciot_paragraphs():
        add_body(document, paragraph)
    add_heading(document, "4.2 CICDDoS2019 Dataset", level=2)
    for paragraph in cicddos_paragraphs():
        add_body(document, paragraph)
    add_heading(document, "4.3 TON-IoT Dataset", level=2)
    for paragraph in toniot_paragraphs():
        add_body(document, paragraph)
    add_table(document, dataset_table, "Table 1. Comparative statistics of the analyzed IoT DDoS datasets.")

    add_heading(document, "Dataset Analysis")
    for paragraph in dataset_analysis_paragraphs():
        add_body(document, paragraph)
    add_figure(document, ANALYSIS_DIR / "paper_figures" / "dataset_size_comparison.png", "Figure 1. Dataset size comparison across IoT DDoS datasets.")
    add_figure(document, ANALYSIS_DIR / "paper_figures" / "class_imbalance_comparison.png", "Figure 2. Class imbalance comparison across IoT DDoS datasets.")

    add_heading(document, "Feature Analysis")
    for paragraph in feature_analysis_paragraphs():
        add_body(document, paragraph)
    add_table(document, redundancy_table, "Table 2. Representative highly redundant feature pairs across the analyzed datasets.")

    add_heading(document, "Experimental Results")
    for paragraph in results_paragraphs():
        add_body(document, paragraph)
    add_figure(document, ADVANCED_DIR / "pca_CICIoT2023.png", "Figure 3. PCA projection of CICIoT2023 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "pca_CICDDoS2019.png", "Figure 4. PCA projection of CICDDoS2019 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "pca_TONIoT.png", "Figure 5. PCA projection of TON-IoT benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "tsne_CICIoT2023.png", "Figure 6. t-SNE projection of CICIoT2023 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "tsne_CICDDoS2019.png", "Figure 7. t-SNE projection of CICDDoS2019 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "tsne_TONIoT.png", "Figure 8. t-SNE projection of TON-IoT benign and DDoS traffic.")
    add_table(document, difficulty_table, "Table 3. Baseline ROC-AUC comparison for Logistic Regression and Random Forest.")
    add_figure(document, ANALYSIS_DIR / "paper_figures" / "dataset_difficulty_comparison.png", "Figure 9. Dataset difficulty comparison using baseline ROC-AUC scores.")

    add_heading(document, "Discussion")
    for paragraph in discussion_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Limitations of Current IoT DDoS Datasets")
    for paragraph in limitations_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Conclusion")
    for paragraph in conclusion_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Future Work")
    for paragraph in future_work_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "References")
    for ref in references():
        add_body(document, ref)

    document.save(str(OUTPUT_PATH))
    print(f"Final paper draft saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
