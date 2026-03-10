from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE_DIR = Path(r"D:\ddos dataset")
ANALYSIS_DIR = BASE_DIR / "analysis_results"
ADVANCED_DIR = ANALYSIS_DIR / "advanced_analysis"
OUTPUT_PATH = BASE_DIR / "iot_ddos_dataset_analysis_paper_v4_expanded.docx"

TITLE = "Comparative Analysis of Modern IoT DDoS Datasets: CICIoT2023, CICDDoS2019, and TON-IoT"
AUTHOR_NAME = "Vinay T. Patil"
AFFILIATION_1 = "Department of Computer Engineering"
AFFILIATION_2 = "Kavayitri Bahinabai Chaudhari North Maharashtra University"
AUTHOR_EMAIL = "Email: vinay@example.com"


def load_tables() -> dict:
    return {
        "dataset_comparison": pd.read_csv(ANALYSIS_DIR / "dataset_comparison_table.csv"),
        "dataset_difficulty": pd.read_csv(ANALYSIS_DIR / "dataset_difficulty_table.csv"),
        "feature_redundancy": pd.read_csv(ANALYSIS_DIR / "feature_redundancy_summary.csv"),
    }


def set_default_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(11)


def set_margins(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_title_block(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in [AUTHOR_NAME, AFFILIATION_1, AFFILIATION_2, AUTHOR_EMAIL]:
        block = document.add_paragraph()
        block.alignment = WD_ALIGN_PARAGRAPH.CENTER
        block.add_run(line)

    document.add_paragraph("")


def add_heading(document: Document, text: str, level: int = 1) -> None:
    p = document.add_paragraph()
    p.style = document.styles[f"Heading {level}"]
    run = p.add_run(text)
    run.bold = True


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)


def add_keywords(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    lead = p.add_run("Keywords: ")
    lead.bold = True
    p.add_run(text)


def add_table(document: Document, dataframe: pd.DataFrame, caption: str) -> None:
    cp = document.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = cp.add_run(caption)
    cap.bold = True

    table = document.add_table(rows=1, cols=len(dataframe.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = table.rows[0].cells
    for i, col in enumerate(dataframe.columns):
        headers[i].text = str(col)
        shade_cell(headers[i], "D9EAF7")
        for p in headers[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True

    for _, row in dataframe.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if isinstance(value, float):
                text = "inf" if value == float("inf") else f"{value:.6f}" if abs(value) < 1000 else f"{value:.2f}"
            else:
                text = str(value)
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph("")


def add_figure(document: Document, image_path: Path, caption: str, width: float = 5.8) -> None:
    if not image_path.exists():
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))

    cp = document.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = cp.add_run(caption)
    cap.italic = True


def references() -> list[str]:
    return [
        "[1] E. C. P. Neto, S. Dadkhah, R. Ferreira, A. Zohourian, R. Lu, and A. A. Ghorbani, \"CICIoT2023: A real-time dataset and benchmark for large-scale attacks in IoT environment,\" Sensors, vol. 23, no. 13, Art. no. 5941, 2023.",
        "[2] I. Sharafaldin, A. H. Lashkari, S. Hakak, and A. A. Ghorbani, \"Developing realistic distributed denial of service (DDoS) attack dataset and taxonomy,\" in Proc. IEEE 53rd Int. Carnahan Conf. Security Technology (ICCST), 2019, pp. 1-8.",
        "[3] A. Alsaedi, N. Moustafa, Z. Tari, A. Mahmood, and A. Anwar, \"TON_IoT telemetry dataset: A new generation dataset of IoT and IIoT for data-driven intrusion detection systems,\" IEEE Access, vol. 8, pp. 165130-165150, 2020.",
        "[4] N. Moustafa, \"A new distributed architecture for evaluating AI-based security systems at the edge: Network TON_IoT datasets,\" Sustainable Cities and Society, vol. 72, Art. no. 102994, 2021.",
        "[5] T. M. Booij, I. Chiscop, E. Meeuwissen, N. Moustafa, and F. T. H. den Hartog, \"ToN_IoT: The role of heterogeneity and the need for standardization of features and attack types in IoT intrusion datasets,\" IEEE Internet of Things Journal, vol. 9, no. 1, pp. 485-496, 2022.",
        "[6] I. Sharafaldin, A. H. Lashkari, and A. A. Ghorbani, \"Toward generating a new intrusion detection dataset and intrusion traffic characterization,\" in Proc. 4th Int. Conf. Information Systems Security and Privacy (ICISSP), 2018, pp. 108-116.",
        "[7] N. Moustafa and J. Slay, \"UNSW-NB15: A comprehensive data set for network intrusion detection systems,\" in Proc. Military Communications and Information Systems Conf. (MilCIS), 2015, pp. 1-6.",
        "[8] N. Moustafa and J. Slay, \"The evaluation of network anomaly detection systems: Statistical analysis of the UNSW-NB15 data set and the comparison with the KDD99 data set,\" Information Security Journal: A Global Perspective, vol. 25, no. 1-3, pp. 18-31, 2016.",
        "[9] A. Khraisat, I. Gondal, P. Vamplew, and J. Kamruzzaman, \"Survey of intrusion detection systems: Techniques, datasets and challenges,\" Cybersecurity, vol. 2, Art. no. 20, 2019.",
        "[10] A. Adnan, A. Muhammed, A. A. A. Ghani, A. Abdullah, and F. Hakim, \"An intrusion detection system for the Internet of Things based on machine learning: Review and challenges,\" Symmetry, vol. 13, no. 6, Art. no. 1011, 2021.",
        "[11] C. Kolias, G. Kambourakis, A. Stavrou, and J. Voas, \"DDoS in the IoT: Mirai and other botnets,\" Computer, vol. 50, no. 7, pp. 80-84, 2017.",
        "[12] S. Garcia, A. Parmisano, and M. J. Erquiaga, \"IoT-23: A labeled dataset with malicious and benign IoT network traffic,\" Zenodo, 2020.",
        "[13] KDD Cup 1999 Data, UCI KDD Archive, Univ. California, Irvine, 1999.",
        "[14] M. Tavallaee, E. Bagheri, W. Lu, and A. A. Ghorbani, \"A detailed analysis of the KDD CUP 99 data set,\" in Proc. IEEE Symp. Computational Intelligence for Security and Defense Applications, 2009, pp. 1-6.",
        "[15] H. Hotelling, \"Analysis of a complex of statistical variables into principal components,\" Journal of Educational Psychology, vol. 24, no. 6, pp. 417-441, 1933.",
        "[16] L. van der Maaten and G. Hinton, \"Visualizing data using t-SNE,\" Journal of Machine Learning Research, vol. 9, pp. 2579-2605, 2008.",
        "[17] L. Breiman, \"Random forests,\" Machine Learning, vol. 45, no. 1, pp. 5-32, 2001.",
        "[18] D. R. Cox, \"The regression analysis of binary sequences,\" Journal of the Royal Statistical Society: Series B, vol. 20, no. 2, pp. 215-242, 1958.",
        "[19] B. R. Kikissagbe and M. Adda, \"Machine learning-based intrusion detection methods in IoT systems: A comprehensive review,\" Electronics, vol. 13, no. 18, Art. no. 3601, 2024.",
        "[20] N. Nisha, N. S. Gill, and P. Gulia, \"A review on machine learning based intrusion detection system for internet of things enabled environment,\" International Journal of Electrical and Computer Engineering, vol. 14, no. 2, pp. 1890-1898, 2024.",
    ]


def abstract_paragraphs() -> list[str]:
    return [
        "The growth of Internet of Things (IoT) infrastructures has intensified the operational importance of reliable distributed denial-of-service (DDoS) detection mechanisms in heterogeneous cyber-physical environments [1]-[5], [9]-[11].",
        "This paper presents a comparative analysis of three contemporary IoT DDoS datasets, namely CICIoT2023, CICDDoS2019, and TON-IoT, using a unified analytical workflow that combines dataset profiling, feature analysis, dimensionality-reduction visualization, and baseline classifier evaluation [1]-[5].",
        "The results indicate that CICIoT2023 and CICDDoS2019 are highly separable under baseline classifiers, whereas TON-IoT presents a comparatively more challenging decision boundary for linear models despite strong nonlinear separability.",
        "The study also identifies severe class imbalance, especially in CICDDoS2019, and strong redundancy among several packet-count and packet-length feature families. These findings provide practical guidance for benchmark selection, feature pruning, and future cross-dataset IoT DDoS detection research."
    ]


def introduction_paragraphs() -> list[str]:
    return [
        "The expansion of IoT technologies has transformed contemporary digital ecosystems by embedding sensors, actuators, gateways, and autonomous devices into homes, hospitals, transportation systems, factories, and urban infrastructure [9], [10], [19], [20]. In parallel with conventional networked systems, these deployments are now tightly integrated with cyber-physical operations in which communication delays, service interruptions, or control failures can have direct operational consequences. Consequently, the security of IoT environments has become inseparable from the resilience of broader cyber-physical systems.",
        "IoT devices play a central role in smart city services, environmental monitoring, connected healthcare, industrial automation, intelligent logistics, and edge-assisted decision support. However, the same properties that enable wide deployment, including low-cost hardware, lightweight operating environments, remote accessibility, and large-scale heterogeneity, also create favorable conditions for adversaries. Weak device authentication, insecure default configurations, and delayed software maintenance continue to increase the exposure of IoT systems to botnet propagation and volumetric service disruption [10], [11], [19].",
        "Among the most damaging attack classes in this context are DDoS attacks, which can exploit compromised IoT devices to exhaust bandwidth, overload servers, or disrupt service availability at scale. The Mirai botnet demonstrated how rapidly poorly secured IoT devices can be weaponized into a high-impact attack infrastructure [11]. Since then, IoT-oriented DDoS activity has remained a major concern for network defenders, particularly in settings where availability and low latency are essential to service delivery.",
        "Traditional intrusion detection systems (IDSs) have struggled to adapt to IoT environments because signature-based methods often fail against evolving attack behavior, while rule-based approaches can be difficult to maintain across highly diverse device populations [9], [10], [19]. Machine-learning-based IDS strategies have therefore gained prominence due to their ability to model complex traffic behavior, support anomaly-aware detection, and generalize across broader attack patterns. Even so, model performance is strongly influenced by the quality and structure of the benchmark datasets used during training and evaluation.",
        "For this reason, benchmark datasets occupy a critical role in IoT security research. Datasets define the attack families, traffic realism, class distributions, and feature spaces from which empirical conclusions are drawn [1]-[8], [12]-[14]. A model evaluated on a highly separable, strongly imbalanced benchmark may not generalize to a more heterogeneous or realistic setting. This study therefore adopts a dataset-centric perspective and analyzes CICIoT2023, CICDDoS2019, and TON-IoT in terms of scale, feature characteristics, separability, redundancy, and practical difficulty in order to support more rigorous DDoS detection research."
    ]


def related_work_paragraphs() -> list[str]:
    return [
        "The evolution of intrusion detection datasets has mirrored the broader transition from traditional enterprise networks to dynamic, heterogeneous cyber environments. Early benchmarks such as KDD Cup 1999 and later re-examinations of that corpus played an important historical role in IDS evaluation, yet they were eventually criticized for outdated attack distributions, duplicated records, and limited realism [13], [14]. These limitations motivated a shift toward more representative datasets that captured richer network behavior and more recent threat scenarios.",
        "Subsequent benchmarks such as UNSW-NB15 and CICIDS2017 aimed to improve realism by combining modern attack categories with more naturalistic traffic generation and richer flow-level feature extraction [6]-[8]. These datasets helped move intrusion detection research beyond legacy corpora by supporting more robust comparisons of anomaly detection, supervised classification, and hybrid methods. Nevertheless, they were still not designed specifically for the heterogeneity and operational constraints of IoT ecosystems.",
        "The emergence of IoT-focused datasets addressed that gap by introducing telemetry data, embedded-device traffic, IoT-specific attack behaviors, and industrial control contexts. TON-IoT contributed a heterogeneous benchmark spanning telemetry, operating system traces, and network traffic across IoT and IIoT scenarios [3]-[5]. IoT-23 similarly emphasized realistic IoT traffic with malicious and benign traces [12]. More recently, CICIoT2023 was introduced as a large-scale IoT benchmark with extensive attack coverage and flow-based features derived from a real-device topology [1].",
        "In parallel, DDoS detection research has increasingly adopted machine-learning-based approaches, including linear classifiers, tree ensembles, deep learning models, and hybrid frameworks, to cope with evolving threat behavior and large feature spaces [9], [10], [19], [20]. These methods often report strong performance on selected benchmarks, but comparisons across studies remain difficult because datasets differ substantially in feature definitions, traffic realism, attack scope, and label design. The literature has therefore highlighted the need for greater standardization and clearer discussion of benchmark properties, especially in IoT settings [5], [9], [19].",
        "Taken together, prior work suggests that the choice of dataset is not a neutral implementation detail but a decisive experimental factor. Modern IoT datasets differ not only in attack categories but also in heterogeneity, telemetry richness, and class structure [1]-[5], [12]. The present study extends this discussion by comparatively analyzing three modern benchmarks from a common methodological perspective, thereby focusing on the datasets themselves rather than proposing a new detection architecture."
    ]


def methodology_paragraphs() -> list[str]:
    return [
        "The methodological pipeline began with dataset acquisition from locally stored processed outputs derived from CICIoT2023, CICDDoS2019, and TON-IoT. These processed versions were created after filtering the original corpora for benign and DDoS traffic relevant to the present binary classification study [1]-[5]. The resulting datasets were standardized into a common label space in which benign traffic was assigned label 0 and DDoS traffic label 1, thereby enabling direct statistical and comparative analysis across otherwise heterogeneous sources.",
        "A dedicated preprocessing workflow was then applied to prepare the feature spaces for analysis. Non-numeric identifiers such as IP addresses, timestamps, and other explicit identifiers were removed, while remaining variables were coerced into numeric form wherever appropriate. Missing values and extreme numerical values were handled through controlled cleaning procedures so that descriptive statistics, visualization, and baseline classifiers could be computed on a consistent representation of each dataset. Because the datasets differ in schema, the analysis focused on cleaned per-dataset feature spaces rather than imposing an artificial global feature alignment.",
        "Feature-oriented analysis was used to quantify multiple structural properties of the cleaned datasets. Redundancy detection relied on correlation analysis to identify highly correlated variable pairs, while entropy analysis was used to measure the informational variability of individual features after discretization. These steps were intended to reveal whether apparent feature richness translated into genuinely diverse information or instead reflected overlapping engineered attributes. Such analysis is important because highly redundant variables may distort interpretation of model importance or inflate effective dimensionality [8]-[10].",
        "To inspect class separability in lower-dimensional space, the study applied principal component analysis (PCA) and t-distributed stochastic neighbor embedding (t-SNE) to sampled and standardized feature matrices [15], [16]. PCA was used as a linear projection technique for identifying dominant variance structure, while t-SNE was employed as a nonlinear visualization method to reveal local clustering behavior that may not be apparent under linear projection alone. For computational tractability, very large datasets were sampled prior to these visualization steps.",
        "Finally, baseline machine-learning experiments were conducted using Logistic Regression and Random Forest classifiers in order to estimate dataset difficulty under both linear and nonlinear decision boundaries [17], [18]. Stratified train-test splits were used to preserve class proportions. These experiments were not intended to propose a new detection model; rather, they served as operational probes of benchmark separability and practical learning complexity. In combination with the statistical and visualization analyses, the model results provide a multi-perspective characterization of dataset difficulty for IoT DDoS research."
    ]


def dataset_description_intro() -> str:
    return (
        "The three datasets analyzed in this study represent different design philosophies in modern intrusion detection research. "
        "CICIoT2023 emphasizes large-scale IoT attack realism, CICDDoS2019 focuses on realistic DDoS taxonomy and flow-based "
        "evaluation, and TON-IoT emphasizes heterogeneous data collection across IoT, IIoT, operating-system, and network layers "
        "[1]-[5]. The local documentation files provided with each dataset further clarify their collection environments, feature "
        "generation strategies, and attack coverage, allowing a more grounded comparative description."
    )


def ciciot_subsection_paragraphs() -> list[str]:
    return [
        "CICIoT2023 was designed as a large-scale benchmark for IoT security analytics and was collected in an extensive IoT topology composed of 105 devices, with attacks launched by malicious IoT devices against other IoT devices [1]. According to the accompanying documentation, the principal objective of the dataset was to support realistic IoT attack analysis by moving beyond small simulated setups and limited-device experiments that had constrained previous research.",
        "The documentation describes the Canadian Institute for Cybersecurity (CIC) IoT laboratory as a dedicated environment created to support realistic IoT security experimentation. The lab was motivated by the difficulty of constructing large, physically realistic IoT topologies using real hardware, switching infrastructure, monitoring equipment, and sustained operational support. This background is important because it contextualizes CICIoT2023 not simply as a large dataset, but as a benchmark collected from a purpose-built real-device environment rather than a purely synthetic simulation.",
        "CICIoT2023 includes 33 attacks categorized into seven high-level classes: DDoS, DoS, Recon, Web-based, Brute Force, Spoofing, and Mirai [1]. Within the DDoS category, the local dataset collection used in this study contains multiple attack families, including ACK fragmentation, ICMP flood, ICMP fragmentation, PSH-ACK flood, RST/FIN flood, SYN flood, SlowLoris, SynonymousIP flood, TCP flood, UDP flood, UDP fragmentation, and HTTP flood. The breadth of these attack families makes the dataset particularly relevant for evaluating whether DDoS detectors generalize across different volumetric and protocol-level attack behaviors.",
        "From a data-format perspective, the documentation specifies that the overall CICIoT2023 distribution contains raw packet captures, extracted CSV features for machine-learning analysis, an example notebook, and supplementary materials describing the collection and wrangling toolchain. The source materials indicate the use of Mergecap, PySpark, TCPDump, and DPKT to manage packet captures and extract flow-based features. In the cleaned subset used here, the retained feature space contains 39 numeric variables, reflecting a compact but highly attack-relevant representation of network behavior."
    ]


def cicddos_subsection_paragraphs() -> list[str]:
    return [
        "CICDDoS2019 was developed to provide a realistic benchmark and taxonomy for distributed denial-of-service attacks in network environments [2]. The documentation emphasizes two linked goals: first, to address shortcomings in prior DDoS datasets, and second, to support both attack detection and family-level classification through rich network-flow features. This makes CICDDoS2019 not only a detection dataset, but also a benchmark explicitly grounded in DDoS taxonomy design.",
        "The dataset documentation distinguishes between reflection-based and exploitation-based DDoS attacks. Reflection-based attacks conceal the attacker behind third-party reflectors and include application-layer protocols such as MSSQL, SSDP, DNS, LDAP, NetBIOS, and SNMP over TCP or UDP. Exploitation-based attacks include SYN flood, UDP flood, and UDP-Lag, reflecting attacks that directly consume bandwidth or protocol state. This taxonomy is directly reflected in the dataset’s attack inventory and remains one of the defining contributions of CICDDoS2019 [2].",
        "A notable element of the collection methodology is the generation of realistic background traffic. The documentation reports that benign background traffic was produced using the B-Profile system, with the abstract behavior of 25 users modeled across HTTP, HTTPS, FTP, SSH, and email protocols. The resulting benchmark includes both benign traffic and modern reflective DDoS attacks such as PortMap, NetBIOS, LDAP, MSSQL, UDP, UDP-Lag, SYN, NTP, DNS, SNMP, SSDP, WebDDoS, and TFTP. This design choice reinforces the dataset’s role as a flow-based evaluation environment that attempts to emulate realistic mixed traffic conditions.",
        "CICDDoS2019 further includes raw PCAPs, event logs, and CSV files derived from CICFlowMeter-V3. The documentation states that more than 80 traffic features were extracted and stored as CSV files, which is consistent with the 82 cleaned numeric features retained in the present binary subset. Because of its richer feature space and explicit DDoS family coverage, CICDDoS2019 remains one of the most useful benchmarks for evaluating flow-level DDoS detection and family differentiation under realistic attack taxonomies [2]."
    ]


def toniot_subsection_paragraphs() -> list[str]:
    return [
        "TON-IoT was introduced as a new generation of IoT and IIoT datasets designed for evaluating artificial-intelligence-based cybersecurity applications in heterogeneous environments [3]-[5]. The accompanying documentation describes the dataset as originating from a realistic and large-scale network deployed at the Cyber Range and IoT Labs of UNSW Canberra at ADFA. This testbed was built to support integrated IoT, cloud, and edge/fog interactions across multiple layers of an Industry 4.0 environment.",
        "The documentation highlights the distributed architecture of the testbed, which used multiple virtual machines and hosts running Windows, Linux, and Kali Linux to manage interconnections between IoT, cloud, and edge/fog systems. More than ten IoT and IIoT sensors, including weather and Modbus-related devices, were used to collect telemetry data. In addition to telemetry, the broader TON-IoT distribution includes network traffic, Windows audit traces, Linux traces, and ground-truth security events, thereby making the dataset inherently heterogeneous in both source type and feature semantics [3]-[5].",
        "Attack scenarios in TON-IoT include DoS, DDoS, ransomware, and other malicious activities directed at web applications, IoT gateways, MQTT-related services, Linux and Windows systems, and network services. The documentation explains that the data were collected in parallel so that normal behavior and cyberattack events could be observed simultaneously across network traffic, operating-system traces, and telemetry streams. This is a significant difference from more narrowly flow-oriented datasets, because it frames TON-IoT as a multi-modal benchmark rather than a single-source traffic corpus.",
        "The dataset structure also reflects this heterogeneity. TON-IoT contains raw datasets, processed datasets, train-test subsets, feature-description and statistics folders, and security-event ground truth. Network data were collected in PCAP, log, and Zeek/Bro CSV formats; Linux traces were recorded using tracing tools such as atop; Windows data were gathered through Performance Monitor; and IoT telemetry was logged in CSV and log files. In the cleaned binary subset used in this study, TON-IoT contributes 19 numeric features, but the broader dataset family remains much richer in modality and attack context than the simplified binary representation analyzed here [3]-[5]."
    ]


def dataset_analysis_paragraphs() -> list[str]:
    return [
        "The comparative statistics reveal that the datasets differ sharply in total volume, benign availability, and feature count. CICIoT2023 contributes the largest cleaned sample size, while CICDDoS2019 provides the most extensive numeric feature space. TON-IoT is smaller in size but exhibits a more moderate DDoS-to-benign ratio than the CIC datasets. These contrasts are not merely descriptive; they have direct implications for how classifier performance should be interpreted.",
        "Class imbalance is particularly severe in CICDDoS2019, where the benign class is extremely limited relative to DDoS traffic. Such imbalance can make high aggregate scores appear easier to achieve even when generalization to realistic benign diversity remains uncertain. CICIoT2023 is also strongly skewed toward attack traffic, although less extremely. TON-IoT, by contrast, provides a relatively less imbalanced scenario and therefore serves as a useful complement when evaluating whether models remain effective under less trivial class conditions.",
        "From a feature perspective, the datasets also differ in how densely engineered their representations appear. CICDDoS2019 includes a rich set of flow-oriented attributes, whereas CICIoT2023 uses a more compact numeric representation tailored to IoT attack traffic. TON-IoT’s cleaned subset is the smallest of the three in terms of numeric features, but this compactness should be viewed in light of the broader heterogeneity of the original benchmark, which spans telemetry, operating system traces, and network data beyond the subset used here.",
        "These dataset-level differences motivate the remainder of the analysis, including feature redundancy, entropy, reduced-dimensional visualization, and baseline model comparison. Together, these steps help determine whether performance differences arise from modeling choices or from the underlying structure of the benchmarks themselves."
    ]


def feature_analysis_paragraphs() -> list[str]:
    return [
        "Feature-oriented inspection reveals that not all variables contribute equally to dataset diversity. High skewness in several variables indicates that a substantial fraction of observations are concentrated near limited value ranges, with comparatively rare high-magnitude activity dominating the distribution. This pattern is expected in attack-oriented traffic, especially where protocol flags, packet counts, or byte-volume attributes respond strongly to bursty malicious behavior.",
        "Entropy analysis complements this observation by identifying which features exhibit richer value variability after discretization. Features with higher entropy may encode more diverse behavior and can therefore be informative for interpretation and model development, whereas low-entropy features may act more like near-constant indicators under specific attack settings. However, high entropy alone does not guarantee unique information if multiple variables encode similar traffic behavior.",
        "Correlation-based redundancy analysis confirms that several features are strongly coupled. In CICIoT2023, closely related count and flag variables form near-duplicate pairs, while CICDDoS2019 contains multiple packet-length and aggregate-count features with very high pairwise correlation. Such redundancy suggests that model capacity may be spent repeatedly on overlapping signals, thereby motivating dimensionality reduction, feature pruning, or ablation studies in downstream experiments."
    ]


def experimental_results_paragraphs() -> list[str]:
    return [
        "The visualization results and baseline classifier experiments provide a practical view of dataset separability. PCA offers a linear projection of the cleaned feature space and highlights whether the dominant variance directions already separate benign and DDoS traffic. In the present study, the PCA plots for CICIoT2023 and CICDDoS2019 suggest relatively clear separation, whereas TON-IoT shows a less trivial structure.",
        "The t-SNE projections offer a complementary nonlinear perspective by preserving neighborhood relationships more effectively than linear projection. These plots reinforce the observation that all three datasets contain meaningful class structure, but they also suggest that TON-IoT exhibits comparatively more complex local organization. This aligns with the expectation that heterogeneous network and telemetry-related environments may generate more nuanced boundaries between benign and attack traffic.",
        "Baseline model evaluation further clarifies these observations. Logistic Regression and Random Forest both achieve extremely strong ROC-AUC values on CICIoT2023 and CICDDoS2019, indicating that benign and DDoS classes are highly separable within the processed subsets. TON-IoT remains more challenging for Logistic Regression, while Random Forest still performs strongly. This difference is analytically meaningful because it implies a less linear but still learnable boundary in TON-IoT."
    ]


def discussion_paragraphs() -> list[str]:
    return [
        "Several important implications follow from the comparative results. First, benchmark difficulty is not uniform across datasets. Near-perfect baseline performance on CICIoT2023 or CICDDoS2019 should not be interpreted as evidence that a detector will generalize equally well to more heterogeneous or less linearly separable settings such as TON-IoT. This is especially relevant for papers that report high accuracy without analyzing dataset structure in detail.",
        "Second, severe class imbalance remains a central methodological concern. In benchmarks where benign traffic is scarce, high detection performance may partly reflect the ease of identifying overwhelmingly attack-dominated patterns rather than truly robust discrimination. Careful evaluation design, including class-aware metrics and discussion of benign scarcity, is therefore essential.",
        "Third, the presence of strongly redundant features implies that effective model performance may be driven by a relatively small set of repeated signals. This does not invalidate benchmark results, but it does suggest that more informative comparisons should include feature-pruning, stability analysis, or cross-dataset transfer experiments. Taken together, these findings support a more dataset-aware approach to IoT DDoS detection research, where benchmark properties are analyzed alongside model performance rather than treated as background detail."
    ]


def conclusion_paragraphs() -> list[str]:
    return [
        "This work presented an expanded comparative analysis of CICIoT2023, CICDDoS2019, and TON-IoT for IoT DDoS detection research. By combining local dataset documentation with statistical profiling, feature analysis, visualization, and baseline modeling, the study provides a more complete understanding of how benchmark properties shape experimental outcomes.",
        "The results show that the three datasets differ substantially in scale, imbalance, feature richness, redundancy, and practical separability. These differences confirm that benchmark selection should be aligned with the intended research claim, especially when evaluating robustness, generalizability, or deployment realism. More broadly, the study demonstrates the value of treating datasets as research objects in their own right rather than merely as inputs to model development."
    ]


def future_work_paragraphs() -> list[str]:
    return [
        "Future work should extend this analysis toward cross-dataset generalization and transfer learning, particularly to determine whether models trained on one benchmark remain reliable on others with different structural properties. This is especially important in IoT security, where heterogeneity in device types, telemetry modalities, and attack conditions can cause significant domain shift.",
        "Additional directions include deep learning and sequence-aware modeling, federated evaluation across distributed IoT environments, and the development of new real-world benchmarks with richer benign traffic diversity. Such efforts would help move IoT DDoS research beyond highly separable benchmark conditions and closer to operationally realistic deployment scenarios."
    ]


def main() -> None:
    data = load_tables()
    dataset_table = data["dataset_comparison"]
    difficulty_table = data["dataset_difficulty"]
    redundancy_table = data["feature_redundancy"].head(12)

    document = Document()
    set_default_font(document)
    set_margins(document)
    add_page_number(document.sections[0])

    add_title_block(document)

    add_heading(document, "Abstract")
    for paragraph in abstract_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Keywords")
    add_keywords(document, "IoT security, DDoS detection, dataset analysis, intrusion detection, machine learning")

    add_heading(document, "Introduction")
    for paragraph in introduction_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Related Work")
    for paragraph in related_work_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Methodology")
    for paragraph in methodology_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Dataset Description")
    add_body(document, dataset_description_intro())

    add_heading(document, "4.1 CICIoT2023 Dataset", level=2)
    for paragraph in ciciot_subsection_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "4.2 CICDDoS2019 Dataset", level=2)
    for paragraph in cicddos_subsection_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "4.3 TON-IoT Dataset", level=2)
    for paragraph in toniot_subsection_paragraphs():
        add_body(document, paragraph)

    add_table(document, dataset_table, "Table 1. Comparative statistics of the analyzed IoT DDoS datasets.")

    add_heading(document, "Dataset Analysis")
    for paragraph in dataset_analysis_paragraphs():
        add_body(document, paragraph)
    add_figure(document, ANALYSIS_DIR / "paper_figures" / "dataset_size_comparison.png", "Figure 1. Dataset size comparison across IoT DDoS datasets.")
    add_figure(document, ANALYSIS_DIR / "paper_figures" / "class_imbalance_comparison.png", "Figure 2. Class imbalance comparison across IoT DDoS datasets.")

    add_heading(document, "Feature Analysis")
    for paragraph in feature_analysis_paragraphs():
        add_body(document, paragraph)
    add_table(document, redundancy_table, "Table 2. Representative highly redundant feature pairs across the analyzed datasets.")

    add_heading(document, "Experimental Results")
    for paragraph in experimental_results_paragraphs():
        add_body(document, paragraph)
    add_figure(document, ADVANCED_DIR / "pca_CICIoT2023.png", "Figure 3. PCA projection of CICIoT2023 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "pca_CICDDoS2019.png", "Figure 4. PCA projection of CICDDoS2019 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "pca_TONIoT.png", "Figure 5. PCA projection of TON-IoT benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "tsne_CICIoT2023.png", "Figure 6. t-SNE projection of CICIoT2023 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "tsne_CICDDoS2019.png", "Figure 7. t-SNE projection of CICDDoS2019 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "tsne_TONIoT.png", "Figure 8. t-SNE projection of TON-IoT benign and DDoS traffic.")
    add_table(document, difficulty_table, "Table 3. Baseline ROC-AUC comparison for Logistic Regression and Random Forest.")
    add_figure(document, ANALYSIS_DIR / "paper_figures" / "dataset_difficulty_comparison.png", "Figure 9. Dataset difficulty comparison using baseline ROC-AUC scores.")

    add_heading(document, "Discussion")
    for paragraph in discussion_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Conclusion")
    for paragraph in conclusion_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "Future Work")
    for paragraph in future_work_paragraphs():
        add_body(document, paragraph)

    add_heading(document, "References")
    for ref in references():
        add_body(document, ref)

    document.save(str(OUTPUT_PATH))
    print(f"Expanded paper draft saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
