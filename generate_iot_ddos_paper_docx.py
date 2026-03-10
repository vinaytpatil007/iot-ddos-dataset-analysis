import os
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE_DIR = Path(r"D:\ddos dataset")
ANALYSIS_DIR = BASE_DIR / "analysis_results"
ADVANCED_DIR = ANALYSIS_DIR / "advanced_analysis"
OUTPUT_PATH = BASE_DIR / "iot_ddos_dataset_analysis_paper.docx"

TITLE = "Comparative Analysis of Modern IoT DDoS Datasets: CICIoT2023, CICDDoS2019, and TON-IoT"


def load_data() -> dict:
    return {
        "dataset_comparison": pd.read_csv(ANALYSIS_DIR / "dataset_comparison_table.csv"),
        "dataset_difficulty": pd.read_csv(ANALYSIS_DIR / "dataset_difficulty_table.csv"),
        "feature_redundancy": pd.read_csv(ANALYSIS_DIR / "feature_redundancy_summary.csv"),
        "paper_summary_md": (ANALYSIS_DIR / "paper_results_summary.md").read_text(encoding="utf-8"),
    }


def set_default_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(11)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_title(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Times New Roman"


def add_keywords(document: Document, keywords: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Keywords: ")
    run.bold = True
    paragraph.add_run(keywords)


def add_section_heading(document: Document, title: str, level: int = 1) -> None:
    heading = document.add_heading(level=level)
    run = heading.add_run(title)
    run.bold = True


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(text)


def add_table_from_dataframe(document: Document, dataframe: pd.DataFrame, title: str) -> None:
    add_paragraph(document, title)
    table = document.add_table(rows=1, cols=len(dataframe.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for idx, column_name in enumerate(dataframe.columns):
        hdr_cells[idx].text = str(column_name)
        set_cell_shading(hdr_cells[idx], "D9E2F3")
        for paragraph in hdr_cells[idx].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True

    for _, row in dataframe.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            if isinstance(value, float):
                if value == float("inf"):
                    text = "inf"
                else:
                    text = f"{value:.6f}" if abs(value) < 1000 else f"{value:.2f}"
            else:
                text = str(value)
            cells[idx].text = text
            for paragraph in cells[idx].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph("")


def add_figure(document: Document, image_path: Path, caption: str, width_inches: float = 5.8) -> None:
    if not image_path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))

    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    caption_run.italic = True


def generate_abstract() -> str:
    return (
        "The rapid expansion of Internet of Things (IoT) deployments has intensified concerns about large-scale "
        "distributed denial-of-service (DDoS) attacks targeting smart environments, edge devices, and connected "
        "services. Reliable intrusion-detection research depends on representative benchmark datasets, yet many "
        "studies still evaluate models on single datasets without examining differences in scale, feature space, "
        "class imbalance, or learning difficulty. This paper presents a comparative analysis of three modern IoT "
        "DDoS datasets: CICIoT2023, CICDDoS2019, and TON-IoT. The study combines statistical profiling, feature "
        "inspection, redundancy analysis, dimensionality-reduction visualization, and baseline machine-learning "
        "evaluation to characterize dataset behavior in a unified manner. The results show that CICIoT2023 and "
        "CICDDoS2019 are highly separable under baseline models, while TON-IoT remains more challenging for linear "
        "classification despite strong nonlinear performance. The analysis also highlights severe class imbalance in "
        "CICDDoS2019 and strong redundancy among several packet-count and packet-length features. These findings "
        "provide practical guidance for dataset selection, feature pruning, and future IoT DDoS detection studies."
    )


def generate_introduction() -> str:
    return (
        "IoT ecosystems now span smart homes, industrial monitoring, healthcare platforms, and urban sensing systems, "
        "creating a dense attack surface composed of heterogeneous, resource-constrained devices. Among the most "
        "damaging threats in this environment are DDoS attacks, which exploit weak device security, default "
        "credentials, and insecure network exposure to overwhelm services or disrupt critical infrastructure. As a "
        "result, data-driven DDoS detection has become a central research topic in IoT security. However, the "
        "quality of any machine-learning-based defense depends heavily on the datasets used for development and "
        "evaluation. Existing studies frequently emphasize model accuracy while underreporting dataset imbalance, "
        "feature redundancy, and cross-dataset difficulty differences. This work addresses that gap by performing a "
        "structured comparative analysis of three widely used modern IoT DDoS datasets. The motivation is not only "
        "to describe dataset statistics, but also to identify how their structural characteristics influence research "
        "claims, reproducibility, and benchmark difficulty."
    )


def generate_related_work() -> str:
    return (
        "Prior research on IoT intrusion-detection datasets has introduced multiple benchmarks designed to capture "
        "network traffic, device telemetry, or mixed attack scenarios. Datasets such as CICDDoS2019 have been widely "
        "used for flow-based DDoS detection, while TON-IoT extends the problem space by including broader cyberattack "
        "families and telemetry sources. More recent datasets such as CICIoT2023 focus explicitly on contemporary IoT "
        "traffic and multiple DDoS variants. In parallel, machine-learning approaches for IoT security have ranged "
        "from linear classifiers and ensemble methods to deep neural architectures. Despite this progress, many "
        "papers evaluate algorithms on a single benchmark and do not sufficiently discuss how dataset composition "
        "affects separability, redundancy, and realism. A comparative dataset-centered perspective is therefore "
        "necessary to interpret reported model performance more rigorously."
    )


def generate_dataset_description_text() -> str:
    return (
        "The study considers three modern IoT DDoS datasets that differ substantially in scale, feature richness, "
        "and class balance. CICIoT2023 provides a large IoT-focused benchmark with multiple DDoS families and a "
        "moderate number of engineered traffic features. CICDDoS2019 offers a richer flow-based feature space and "
        "very large attack volume, but the benign class is extremely limited after filtering. TON-IoT contains a "
        "smaller cleaned subset for binary DDoS analysis and exhibits a less extreme class imbalance than the two CIC "
        "datasets. Table 1 summarizes the resulting dataset statistics used in this comparative study."
    )


def generate_characteristics_text() -> str:
    return (
        "Dataset size and class balance have direct consequences for both training dynamics and evaluation realism. "
        "CICIoT2023 is the largest dataset in this study, while CICDDoS2019 exhibits the most severe DDoS-to-benign "
        "imbalance. TON-IoT is smaller but more balanced, which makes it useful for observing whether model "
        "performance remains strong under less trivial class conditions. Feature counts also vary significantly, with "
        "CICDDoS2019 providing the largest cleaned numeric feature space."
    )


def generate_feature_analysis_text() -> str:
    return (
        "Feature-level analysis reveals that several variables carry high entropy and strong discriminative potential, "
        "while others are nearly deterministic or highly redundant. Redundancy is especially visible in packet-count "
        "and packet-length families, where near-duplicate measurements or tightly coupled transformations appear. "
        "Such redundancy can inflate model complexity without adding meaningful information, making feature pruning an "
        "important step for later experimental design."
    )


def generate_visualization_text() -> str:
    return (
        "PCA and t-SNE visualizations were used to inspect whether benign and DDoS traffic form distinguishable "
        "clusters in reduced-dimensional space. The resulting plots show clear separation in the CIC datasets and a "
        "less trivial, though still structured, separation pattern for TON-IoT. This observation is consistent with "
        "the later classifier-based difficulty results and suggests that dataset geometry differs meaningfully across "
        "benchmarks."
    )


def generate_difficulty_text() -> str:
    return (
        "Baseline model performance provides an operational view of dataset difficulty. CICIoT2023 and CICDDoS2019 "
        "produce near-perfect ROC-AUC values under both Logistic Regression and Random Forest baselines, indicating "
        "that benign and DDoS classes are highly separable in the sampled experimental setting. TON-IoT remains more "
        "challenging for Logistic Regression, although Random Forest still performs strongly. This gap suggests a more "
        "nonlinear decision boundary and makes TON-IoT particularly relevant for robust DDoS detection research."
    )


def generate_discussion_text() -> str:
    return (
        "The comparative results point to three major issues for IoT DDoS research. First, class imbalance can distort "
        "reported performance, especially when benign traffic is scarce, as observed in CICDDoS2019. Second, feature "
        "redundancy may cause models to appear stronger than necessary if highly correlated variables dominate the "
        "decision process. Third, dataset difficulty is not uniform: high benchmark scores on CICIoT2023 or "
        "CICDDoS2019 do not imply equivalent robustness on TON-IoT or on unseen real-world IoT traffic. Therefore, "
        "benchmark selection should be tied explicitly to the intended deployment scenario and research claim."
    )


def generate_conclusion_text() -> str:
    return (
        "This paper presented a comparative analysis of CICIoT2023, CICDDoS2019, and TON-IoT for IoT DDoS detection "
        "research. The study combined dataset statistics, feature inspection, redundancy analysis, dimensionality "
        "reduction, and baseline model evaluation to characterize benchmark behavior. The results demonstrate major "
        "differences in class balance, feature composition, redundancy, and practical difficulty. These findings "
        "support more transparent dataset selection and create a stronger foundation for future benchmark-driven IoT "
        "DDoS studies."
    )


def generate_future_work_text() -> str:
    return (
        "Future work should investigate cross-dataset generalization to determine whether models trained on one IoT "
        "benchmark transfer effectively to another. Deep learning architectures, graph-based models, and temporal "
        "sequence methods may also offer additional insight into complex IoT traffic behavior beyond what baseline "
        "classifiers capture. Finally, future dataset development should prioritize more realistic benign traffic, "
        "multi-device diversity, and contemporary real-world IoT attack traces to better support practical DDoS "
        "detection research."
    )


def main() -> None:
    data = load_data()
    dataset_table = data["dataset_comparison"]
    difficulty_table = data["dataset_difficulty"]
    redundancy_table = data["feature_redundancy"]

    document = Document()
    set_default_font(document)

    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    add_title(document, TITLE)
    document.add_paragraph("")

    add_section_heading(document, "Abstract")
    add_paragraph(document, generate_abstract())

    add_section_heading(document, "Keywords")
    add_keywords(document, "IoT security, DDoS detection, dataset analysis, intrusion detection, machine learning")

    add_section_heading(document, "Introduction")
    add_paragraph(document, generate_introduction())

    add_section_heading(document, "Related Work")
    add_paragraph(document, generate_related_work())

    add_section_heading(document, "Dataset Description")
    add_paragraph(document, generate_dataset_description_text())
    add_table_from_dataframe(document, dataset_table, "Table 1. Comparative statistics of the analyzed IoT DDoS datasets.")

    add_section_heading(document, "Dataset Characteristics")
    add_paragraph(document, generate_characteristics_text())
    add_figure(
        document,
        ANALYSIS_DIR / "paper_figures" / "dataset_size_comparison.png",
        "Figure 1: Dataset size comparison across IoT DDoS datasets",
    )
    add_figure(
        document,
        ANALYSIS_DIR / "paper_figures" / "class_imbalance_comparison.png",
        "Figure 2: Class imbalance comparison across IoT DDoS datasets",
    )

    add_section_heading(document, "Feature Analysis")
    add_paragraph(document, generate_feature_analysis_text())
    add_table_from_dataframe(
        document,
        redundancy_table.head(12),
        "Table 2. Highly redundant feature pairs identified across the analyzed datasets.",
    )

    add_section_heading(document, "Dataset Visualization")
    add_paragraph(document, generate_visualization_text())
    add_figure(
        document,
        ADVANCED_DIR / "pca_CICIoT2023.png",
        "Figure 3: PCA projection of CICIoT2023 benign and DDoS traffic",
    )
    add_figure(
        document,
        ADVANCED_DIR / "pca_CICDDoS2019.png",
        "Figure 4: PCA projection of CICDDoS2019 benign and DDoS traffic",
    )
    add_figure(
        document,
        ADVANCED_DIR / "pca_TONIoT.png",
        "Figure 5: PCA projection of TON-IoT benign and DDoS traffic",
    )
    add_figure(
        document,
        ADVANCED_DIR / "tsne_CICIoT2023.png",
        "Figure 6: t-SNE projection of CICIoT2023 benign and DDoS traffic",
    )
    add_figure(
        document,
        ADVANCED_DIR / "tsne_CICDDoS2019.png",
        "Figure 7: t-SNE projection of CICDDoS2019 benign and DDoS traffic",
    )
    add_figure(
        document,
        ADVANCED_DIR / "tsne_TONIoT.png",
        "Figure 8: t-SNE projection of TON-IoT benign and DDoS traffic",
    )

    add_section_heading(document, "Dataset Difficulty Analysis")
    add_paragraph(document, generate_difficulty_text())
    add_table_from_dataframe(document, difficulty_table, "Table 3. ROC-AUC comparison of baseline models across datasets.")
    add_figure(
        document,
        ANALYSIS_DIR / "paper_figures" / "dataset_difficulty_comparison.png",
        "Figure 9: Dataset difficulty comparison using baseline ROC-AUC scores",
    )

    add_section_heading(document, "Discussion")
    add_paragraph(document, generate_discussion_text())

    add_section_heading(document, "Conclusion")
    add_paragraph(document, generate_conclusion_text())

    add_section_heading(document, "Future Work")
    add_paragraph(document, generate_future_work_text())

    document.save(str(OUTPUT_PATH))
    print(f"Paper draft saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
