from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE_DIR = Path(r"D:\ddos dataset")
ANALYSIS_DIR = BASE_DIR / "analysis_results"
ADVANCED_DIR = ANALYSIS_DIR / "advanced_analysis"
OUTPUT_PATH = BASE_DIR / "iot_ddos_dataset_analysis_paper_v3_with_citations.docx"

TITLE = "Comparative Analysis of Modern IoT DDoS Datasets: CICIoT2023, CICDDoS2019, and TON-IoT"
AUTHOR_NAME = "Vinay T. Patil"
AFFILIATION_1 = "Department of Computer Engineering"
AFFILIATION_2 = "Kavayitri Bahinabai Chaudhari North Maharashtra University"
AUTHOR_EMAIL = "Email: vinay@example.com"


def load_tables() -> dict:
    return {
        "dataset_comparison": pd.read_csv(ANALYSIS_DIR / "dataset_comparison_table.csv"),
        "dataset_difficulty": pd.read_csv(ANALYSIS_DIR / "dataset_difficulty_table.csv"),
        "feature_redundancy": pd.read_csv(ANALYSIS_DIR / "feature_redundancy_summary.csv"),
    }


def set_default_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(11)


def set_margins(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_title_block(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in [AUTHOR_NAME, AFFILIATION_1, AFFILIATION_2, AUTHOR_EMAIL]:
        block = document.add_paragraph()
        block.alignment = WD_ALIGN_PARAGRAPH.CENTER
        block.add_run(line)

    document.add_paragraph("")


def add_heading(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.style = document.styles["Heading 1"]
    run = p.add_run(text)
    run.bold = True


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.add_run(text)


def add_keywords(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    lead = p.add_run("Keywords: ")
    lead.bold = True
    p.add_run(text)


def add_table(document: Document, dataframe: pd.DataFrame, caption: str) -> None:
    cp = document.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = cp.add_run(caption)
    cap.bold = True

    table = document.add_table(rows=1, cols=len(dataframe.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = table.rows[0].cells
    for i, col in enumerate(dataframe.columns):
        headers[i].text = str(col)
        shade_cell(headers[i], "D9EAF7")
        for p in headers[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True

    for _, row in dataframe.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if isinstance(value, float):
                text = "inf" if value == float("inf") else f"{value:.6f}" if abs(value) < 1000 else f"{value:.2f}"
            else:
                text = str(value)
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph("")


def add_figure(document: Document, image_path: Path, caption: str, width: float = 5.8) -> None:
    if not image_path.exists():
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))

    cp = document.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = cp.add_run(caption)
    cap.italic = True


def references() -> list[str]:
    return [
        "[1] E. C. P. Neto, S. Dadkhah, R. Ferreira, A. Zohourian, R. Lu, and A. A. Ghorbani, \"CICIoT2023: A real-time dataset and benchmark for large-scale attacks in IoT environment,\" Sensors, vol. 23, no. 13, Art. no. 5941, 2023.",
        "[2] I. Sharafaldin, A. H. Lashkari, S. Hakak, and A. A. Ghorbani, \"Developing realistic distributed denial of service (DDoS) attack dataset and taxonomy,\" in Proc. IEEE 53rd Int. Carnahan Conf. Security Technology (ICCST), 2019, pp. 1-8.",
        "[3] A. Alsaedi, N. Moustafa, Z. Tari, A. Mahmood, and A. Anwar, \"TON_IoT telemetry dataset: A new generation dataset of IoT and IIoT for data-driven intrusion detection systems,\" IEEE Access, vol. 8, pp. 165130-165150, 2020.",
        "[4] N. Moustafa, \"A new distributed architecture for evaluating AI-based security systems at the edge: Network TON_IoT datasets,\" Sustainable Cities and Society, vol. 72, Art. no. 102994, 2021.",
        "[5] T. M. Booij, I. Chiscop, E. Meeuwissen, N. Moustafa, and F. T. H. den Hartog, \"ToN_IoT: The role of heterogeneity and the need for standardization of features and attack types in IoT intrusion datasets,\" IEEE Internet of Things Journal, vol. 9, no. 1, pp. 485-496, 2022.",
        "[6] I. Sharafaldin, A. H. Lashkari, and A. A. Ghorbani, \"Toward generating a new intrusion detection dataset and intrusion traffic characterization,\" in Proc. 4th Int. Conf. Information Systems Security and Privacy (ICISSP), 2018, pp. 108-116.",
        "[7] N. Moustafa and J. Slay, \"UNSW-NB15: A comprehensive data set for network intrusion detection systems,\" in Proc. Military Communications and Information Systems Conf. (MilCIS), 2015, pp. 1-6.",
        "[8] N. Moustafa and J. Slay, \"The evaluation of network anomaly detection systems: Statistical analysis of the UNSW-NB15 data set and the comparison with the KDD99 data set,\" Information Security Journal: A Global Perspective, vol. 25, no. 1-3, pp. 18-31, 2016.",
        "[9] A. Khraisat, I. Gondal, P. Vamplew, and J. Kamruzzaman, \"Survey of intrusion detection systems: Techniques, datasets and challenges,\" Cybersecurity, vol. 2, Art. no. 20, 2019.",
        "[10] A. Adnan, A. Muhammed, A. A. A. Ghani, A. Abdullah, and F. Hakim, \"An intrusion detection system for the Internet of Things based on machine learning: Review and challenges,\" Symmetry, vol. 13, no. 6, Art. no. 1011, 2021.",
        "[11] C. Kolias, G. Kambourakis, A. Stavrou, and J. Voas, \"DDoS in the IoT: Mirai and other botnets,\" Computer, vol. 50, no. 7, pp. 80-84, 2017.",
        "[12] S. Garcia, A. Parmisano, and M. J. Erquiaga, \"IoT-23: A labeled dataset with malicious and benign IoT network traffic,\" Zenodo, ver. 1.0.0, 2020.",
    ]


def abstract_text() -> str:
    return (
        "The proliferation of Internet of Things (IoT) devices has intensified the operational risk posed by distributed "
        "denial-of-service (DDoS) attacks, thereby increasing the importance of high-quality benchmark datasets for "
        "data-driven security research [1]-[5]. This paper presents a comparative analysis of three contemporary IoT "
        "DDoS datasets: CICIoT2023, CICDDoS2019, and TON-IoT [1]-[4]. The study integrates descriptive statistics, "
        "class-imbalance assessment, feature redundancy analysis, dimensionality-reduction visualization, and baseline "
        "machine-learning evaluation to characterize dataset behavior from both data-centric and model-centric "
        "perspectives. The results demonstrate that CICIoT2023 and CICDDoS2019 are highly separable under baseline "
        "classifiers, whereas TON-IoT presents a comparatively more challenging decision boundary for linear models. "
        "The analysis also reveals severe imbalance in CICDDoS2019 and substantial redundancy among packet-count and "
        "packet-length variables across the datasets. These findings underscore the need to interpret reported model "
        "performance in light of dataset structure rather than accuracy alone and provide guidance for benchmark "
        "selection, feature pruning, and cross-dataset evaluation in future IoT DDoS detection studies."
    )


def introduction_text() -> str:
    return (
        "The large-scale deployment of IoT systems in smart homes, industrial automation, healthcare, and urban "
        "infrastructure has created a highly heterogeneous and increasingly vulnerable cyber-physical ecosystem [9], "
        "[10]. Within this environment, DDoS attacks remain one of the most disruptive threats because compromised IoT "
        "devices can be orchestrated into botnets that overwhelm critical services or network resources [11]. As a "
        "result, machine-learning-based intrusion detection has become central to IoT security research [9], [10]. "
        "However, model-centric studies often overlook the influence of dataset composition, feature structure, and "
        "class balance on the credibility of reported detection performance. This issue is especially important in IoT "
        "security, where benchmark datasets differ widely in feature engineering strategies, attack coverage, and "
        "traffic realism [1]-[6]. Accordingly, this work adopts a dataset-centric perspective and examines the relative "
        "characteristics of CICIoT2023, CICDDoS2019, and TON-IoT in order to support more rigorous benchmark selection "
        "and experimental interpretation [1]-[5]."
    )


def related_work_text() -> str:
    return (
        "Benchmark dataset design has long been recognized as a critical issue in network intrusion detection research "
        "[6]-[10]. Earlier efforts such as UNSW-NB15 and CICIDS2017 attempted to address limitations in legacy "
        "benchmarks by improving realism, traffic diversity, and attack representation [6]-[8]. More recent work has "
        "focused specifically on IoT and IIoT environments, leading to datasets such as TON-IoT and IoT-23, which "
        "capture heterogeneous telemetry and realistic malicious behavior [3]-[5], [12]. In parallel, literature on "
        "IoT intrusion detection and DDoS mitigation has emphasized the role of machine learning, dimensionality "
        "reduction, and ensemble methods for distinguishing benign and attack traffic [9]-[11]. Nonetheless, fewer "
        "studies directly compare modern IoT DDoS datasets with respect to separability, redundancy, and benchmark "
        "difficulty. The present study addresses that gap by treating dataset quality itself as an object of "
        "investigation."
    )


def methodology_text() -> str:
    return (
        "The methodology followed a structured analytical pipeline. First, processed binary-class subsets were obtained "
        "from CICIoT2023, CICDDoS2019, and TON-IoT, each of which has been introduced in prior dataset literature [1]-"
        "[4]. Second, preprocessing standardized the attack labels, removed non-numeric identifiers, and retained "
        "cleaned feature matrices suitable for comparative analysis. Third, statistical analysis was used to quantify "
        "dataset size, class imbalance, feature skewness, entropy, and correlation-based redundancy. Fourth, PCA and "
        "t-SNE were applied to sampled and standardized feature spaces to inspect benign-versus-DDoS separation in "
        "reduced-dimensional representations. Finally, Logistic Regression and Random Forest baselines were evaluated "
        "under stratified splits to estimate practical dataset difficulty under linear and nonlinear decision "
        "boundaries, in line with common machine-learning-based IDS evaluation practice [9], [10]."
    )


def dataset_description_text() -> str:
    return (
        "CICIoT2023 provides a recent large-scale IoT attack benchmark built around real IoT devices and multiple "
        "attack families, including extensive DDoS variants [1]. CICDDoS2019 was designed to reflect realistic DDoS "
        "attack scenarios and taxonomy development in flow-based network environments [2]. TON-IoT extends the "
        "benchmark space by providing heterogeneous telemetry and network data collected from IoT and IIoT settings "
        "[3], [4]. In addition, the broader TON-IoT literature has highlighted the importance of feature and label "
        "standardization when comparing heterogeneous IoT intrusion datasets [5]. These properties make the three "
        "datasets well suited for a comparative benchmark analysis."
    )


def dataset_analysis_text() -> str:
    return (
        "The comparative dataset statistics reveal marked differences in scale and class balance. CICIoT2023 is the "
        "largest benchmark among the three, while CICDDoS2019 exhibits the most extreme DDoS-to-benign imbalance. "
        "TON-IoT is smaller but comparatively less skewed, which makes it useful for evaluating whether high model "
        "performance persists under less trivial class conditions. These observations are important because benchmark "
        "imbalance can substantially affect the interpretation of detection metrics, especially in intrusion detection "
        "research that relies heavily on accuracy-oriented reporting [9], [10]."
    )


def feature_analysis_text() -> str:
    return (
        "Feature-level inspection further indicates that multiple variables exhibit high skewness and strong redundancy. "
        "In CICIoT2023, several count-based and flag-based variables are almost perfectly correlated, while "
        "CICDDoS2019 contains multiple packet-length attributes that reflect near-duplicate information. Such behavior "
        "is consistent with concerns in the intrusion-detection literature regarding dataset complexity, high "
        "dimensionality, and the risk of overestimating feature diversity [8]-[10]. Consequently, feature-pruning and "
        "ablation experiments should be considered essential in downstream model studies."
    )


def experimental_results_text() -> str:
    return (
        "Dimensionality-reduction visualizations and baseline model results jointly indicate that the three datasets do "
        "not present equal levels of classification difficulty. The PCA and t-SNE plots show strong benign-versus-DDoS "
        "separation for the CIC datasets, whereas TON-IoT exhibits comparatively less trivial clustering behavior. This "
        "pattern is reinforced by baseline evaluation: Logistic Regression and Random Forest achieve near-perfect "
        "discrimination on CICIoT2023 and CICDDoS2019, while TON-IoT remains more challenging for the linear model. "
        "These findings align with the broader understanding that benchmark-specific structure strongly shapes observed "
        "model performance in IDS research [9]-[11]."
    )


def discussion_text() -> str:
    return (
        "Several implications follow from the comparative analysis. First, severe imbalance, especially in "
        "CICDDoS2019, may inflate confidence in highly accurate detectors if benign scarcity is not explicitly "
        "considered during evaluation [2], [9]. Second, strong redundancy among engineered flow features suggests that "
        "feature selection may reduce complexity without materially degrading detection performance [8]-[10]. Third, "
        "the more challenging linear separability observed for TON-IoT suggests that conclusions drawn from easier "
        "benchmarks cannot be assumed to generalize across heterogeneous IoT environments [3]-[5]. Therefore, DDoS "
        "detection studies should report not only classifier metrics, but also dataset characteristics that contextualize "
        "those results."
    )


def conclusion_text() -> str:
    return (
        "This study presented a comparative analysis of CICIoT2023, CICDDoS2019, and TON-IoT for IoT DDoS detection "
        "research. By integrating statistical profiling, redundancy analysis, reduced-dimensional visualization, and "
        "baseline evaluation, the work demonstrates that benchmark datasets differ substantially in scale, balance, "
        "feature structure, and effective learning difficulty. The findings support more transparent dataset selection "
        "and provide a stronger empirical basis for future IoT intrusion-detection studies [1]-[5], [9], [10]."
    )


def future_work_text() -> str:
    return (
        "Future work should examine cross-dataset generalization, domain shift, and representation learning across "
        "modern IoT DDoS benchmarks [5], [9], [10]. Additional research is also warranted on deep learning and "
        "sequence-aware methods that can exploit temporal structure in IoT traffic, as well as on new real-world "
        "benchmarks that better reflect evolving botnet behavior and heterogeneous edge deployments [4], [11], [12]."
    )


def main() -> None:
    data = load_tables()
    dataset_table = data["dataset_comparison"]
    difficulty_table = data["dataset_difficulty"]
    redundancy_table = data["feature_redundancy"].head(12)

    document = Document()
    set_default_font(document)
    set_margins(document)
    add_page_number(document.sections[0])

    add_title_block(document)

    add_heading(document, "Abstract")
    add_body(document, abstract_text())

    add_heading(document, "Keywords")
    add_keywords(document, "IoT security, DDoS detection, dataset analysis, intrusion detection, machine learning")

    add_heading(document, "Introduction")
    add_body(document, introduction_text())

    add_heading(document, "Related Work")
    add_body(document, related_work_text())

    add_heading(document, "Methodology")
    add_body(document, methodology_text())

    add_heading(document, "Dataset Description")
    add_body(document, dataset_description_text())
    add_table(document, dataset_table, "Table 1. Comparative statistics of the analyzed IoT DDoS datasets.")

    add_heading(document, "Dataset Analysis")
    add_body(document, dataset_analysis_text())
    add_figure(document, ANALYSIS_DIR / "paper_figures" / "dataset_size_comparison.png", "Figure 1. Dataset size comparison across IoT DDoS datasets.")
    add_figure(document, ANALYSIS_DIR / "paper_figures" / "class_imbalance_comparison.png", "Figure 2. Class imbalance comparison across IoT DDoS datasets.")

    add_heading(document, "Feature Analysis")
    add_body(document, feature_analysis_text())
    add_table(document, redundancy_table, "Table 2. Representative highly redundant feature pairs across the analyzed datasets.")

    add_heading(document, "Experimental Results")
    add_body(document, experimental_results_text())
    add_figure(document, ADVANCED_DIR / "pca_CICIoT2023.png", "Figure 3. PCA projection of CICIoT2023 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "pca_CICDDoS2019.png", "Figure 4. PCA projection of CICDDoS2019 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "pca_TONIoT.png", "Figure 5. PCA projection of TON-IoT benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "tsne_CICIoT2023.png", "Figure 6. t-SNE projection of CICIoT2023 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "tsne_CICDDoS2019.png", "Figure 7. t-SNE projection of CICDDoS2019 benign and DDoS traffic.")
    add_figure(document, ADVANCED_DIR / "tsne_TONIoT.png", "Figure 8. t-SNE projection of TON-IoT benign and DDoS traffic.")
    add_table(document, difficulty_table, "Table 3. Baseline ROC-AUC comparison for Logistic Regression and Random Forest.")
    add_figure(document, ANALYSIS_DIR / "paper_figures" / "dataset_difficulty_comparison.png", "Figure 9. Dataset difficulty comparison using baseline ROC-AUC scores.")

    add_heading(document, "Discussion")
    add_body(document, discussion_text())

    add_heading(document, "Conclusion")
    add_body(document, conclusion_text())

    add_heading(document, "Future Work")
    add_body(document, future_work_text())

    add_heading(document, "References")
    for ref in references():
        add_body(document, ref)

    document.save(str(OUTPUT_PATH))
    print(f"Cited paper draft saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
